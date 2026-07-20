"""
Service d'export de conversations IA en PDF et DOCX
avec mise en page professionnelle aux couleurs EDUAI Learning.
Supporte l'export de messages individuels avec détection de langue (Arabe/Français)
et alignement RTL/LTR approprié.
"""

import io
import os
import re
import html as html_lib
from datetime import datetime

_ARABIC_AVAILABLE = False
try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    _ARABIC_AVAILABLE = True
except ImportError:
    arabic_reshaper = None
    get_display = None


_FONT_DIR = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
_ARIAL_REGULAR = os.path.join(_FONT_DIR, "arial.ttf")
_ARIAL_BOLD = os.path.join(_FONT_DIR, "arialbd.ttf")
_ARIAL_ITALIC = os.path.join(_FONT_DIR, "ariali.ttf")
_ARIAL_BOLD_ITALIC = os.path.join(_FONT_DIR, "arialbi.ttf")
_TAHOMA_REGULAR = os.path.join(_FONT_DIR, "tahoma.ttf")
_TAHOMA_BOLD = os.path.join(_FONT_DIR, "tahomabd.ttf")


COLORS = {
    "navy": "#0D1B2A",
    "orange": "#FF6B2B",
    "orange_light": "#FFB347",
    "cream": "#FFF8F0",
    "white": "#FFFFFF",
    "gray": "#6B7280",
    "gray_light": "#F3F4F6",
}


def _format_date(dt: datetime) -> str:
    if not dt:
        return ""
    return dt.strftime("%d/%m/%Y à %H:%M")


def detect_language(text: str) -> str:
    """Détecte si le texte est principalement en arabe ou en français."""
    if not text:
        return "fr"
    arabic_chars = len(re.findall(r'[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]', text))
    latin_chars = len(re.findall(r'[a-zA-ZÀ-ÿ]', text))
    total = arabic_chars + latin_chars
    if total == 0:
        return "fr"
    return "ar" if arabic_chars / total > 0.3 else "fr"


def _reshape_arabic(text: str) -> str:
    """Redessine le texte arabe pour les ligatures et l'ordre bidirectionnel.
    Nécessaire pour fpdf2 qui ne gère pas nativement les ligatures arabes."""
    if not _ARABIC_AVAILABLE or not text:
        return text
    if detect_language(text) != "ar":
        return text
    reshaped = arabic_reshaper.reshape(text)
    return get_display(reshaped)


# ═══════════════════════════════════════════════════════════
#  PDF — Message unique
# ═══════════════════════════════════════════════════════════

def _generate_pdf_from_text(lines: list, title: str = "", is_rtl: bool = False) -> io.BytesIO:
    """Génère un PDF à partir de texte avec support RTL/LTR et police arabe."""
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_font("Arial", "B", 20)
            self.set_text_color(13, 27, 42)
            self.cell(0, 15, "EDUAI Learning", new_x="LMARGIN", new_y="NEXT", align="C")
            self.set_font("Arial", "I", 10)
            self.set_text_color(107, 114, 128)
            self.cell(0, 8, f"Export de reponse IA — {_format_date(datetime.now())}",
                      new_x="LMARGIN", new_y="NEXT", align="C")
            self.set_draw_color(255, 107, 43)
            self.set_line_width(0.5)
            self.line(15, self.get_y() + 2, 195, self.get_y() + 2)
            self.ln(8)

        def footer(self):
            self.set_y(-20)
            self.set_font("Arial", "I", 8)
            self.set_text_color(107, 114, 128)
            self.cell(0, 10, "EDUAI Learning — Plateforme educative intelligente",
                      new_x="LMARGIN", new_y="NEXT", align="C")
            self.cell(0, 5, f"Page {self.page_no()}/{{nb}}", align="C")

    pdf = PDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=25)

    pdf.add_font("Arial", "", _ARIAL_REGULAR)
    pdf.add_font("Arial", "B", _ARIAL_BOLD)
    pdf.add_font("Arial", "I", _ARIAL_ITALIC)
    pdf.add_font("Arial", "BI", _ARIAL_BOLD_ITALIC)
    pdf.add_font("Tahoma", "", _TAHOMA_REGULAR)
    pdf.add_font("Tahoma", "B", _TAHOMA_BOLD)

    pdf.add_page()

    if title:
        pdf.set_font("Arial", "B", 16)
        pdf.set_text_color(13, 27, 42)
        pdf.cell(0, 10, title, new_x="LMARGIN", new_y="NEXT", align="C")
        pdf.ln(5)

    pdf.set_font("Arial", "B", 10)
    pdf.set_fill_color(13, 27, 42)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 8, " Reponse de l'Assistant IA ", new_x="LMARGIN", new_y="NEXT", align="C", fill=True)
    pdf.ln(8)

    align = "R" if is_rtl else "L"
    if is_rtl:
        pdf.set_font("Tahoma", "", 11)
    else:
        pdf.set_font("Arial", "", 11)
    pdf.set_text_color(26, 26, 26)

    for line in lines:
        pdf.set_x(15)
        display_line = _reshape_arabic(line) if is_rtl else line
        pdf.multi_cell(180, 7, display_line, new_x="LMARGIN", new_y="NEXT", align=align)

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return buf


def export_message_pdf(message_content: str, message_role: str = "assistant", title: str = "") -> io.BytesIO:
    """Génère un PDF pour un message unique avec détection de langue RTL/LTR."""
    lang = detect_language(message_content)
    is_rtl = lang == "ar"
    lines = message_content.split("\n")
    display_title = title or "Reponse IA"
    return _generate_pdf_from_text(lines, title=display_title, is_rtl=is_rtl)


# ═══════════════════════════════════════════════════════════
#  PDF — Conversation complète
# ═══════════════════════════════════════════════════════════

def export_pdf(conversation, messages) -> io.BytesIO:
    """Génère un PDF de conversation complète avec RTL/LTR par message."""
    from fpdf import FPDF

    class PDF(FPDF):
        def header(self):
            self.set_font("Arial", "B", 20)
            self.set_text_color(13, 27, 42)
            self.cell(0, 15, "EDUAI Learning", new_x="LMARGIN", new_y="NEXT", align="C")
            self.set_font("Arial", "I", 10)
            self.set_text_color(107, 114, 128)
            self.cell(0, 8, f"Export de conversation — {_format_date(datetime.now())}",
                      new_x="LMARGIN", new_y="NEXT", align="C")
            self.set_draw_color(255, 107, 43)
            self.set_line_width(0.5)
            self.line(15, self.get_y() + 2, 195, self.get_y() + 2)
            self.ln(8)

        def footer(self):
            self.set_y(-20)
            self.set_font("Arial", "I", 8)
            self.set_text_color(107, 114, 128)
            self.cell(0, 10, "EDUAI Learning — Plateforme educative intelligente",
                      new_x="LMARGIN", new_y="NEXT", align="C")
            self.cell(0, 5, f"Page {self.page_no()}/{{nb}}", align="C")

    pdf = PDF()
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=25)

    pdf.add_font("Arial", "", _ARIAL_REGULAR)
    pdf.add_font("Arial", "B", _ARIAL_BOLD)
    pdf.add_font("Arial", "I", _ARIAL_ITALIC)
    pdf.add_font("Arial", "BI", _ARIAL_BOLD_ITALIC)
    pdf.add_font("Tahoma", "", _TAHOMA_REGULAR)
    pdf.add_font("Tahoma", "B", _TAHOMA_BOLD)

    pdf.add_page()

    pdf.set_font("Arial", "B", 16)
    pdf.set_text_color(13, 27, 42)
    pdf.cell(0, 10, conversation.title, new_x="LMARGIN", new_y="NEXT", align="L")

    pdf.set_font("Arial", "", 9)
    pdf.set_text_color(107, 114, 128)
    meta = f"Date : {_format_date(conversation.created_at)}  |  Messages : {len(messages)}"
    pdf.cell(0, 6, meta, new_x="LMARGIN", new_y="NEXT", align="L")
    pdf.ln(5)

    orange_rgb = (255, 107, 43)
    navy_rgb = (13, 27, 42)

    for msg in messages:
        is_user = msg.role == "user"
        role_label = "Eleve" if is_user else "Assistant IA"
        color = orange_rgb if is_user else navy_rgb

        pdf.set_font("Arial", "B", 10)
        pdf.set_text_color(*color)
        pdf.cell(0, 8, f"● {role_label}", new_x="LMARGIN", new_y="NEXT")

        msg_lang = getattr(msg, "detected_language", None) or detect_language(msg.content)
        is_rtl = msg_lang == "ar"
        align = "R" if is_rtl else "L"

        if is_rtl:
            pdf.set_font("Tahoma", "", 11)
        else:
            pdf.set_font("Arial", "", 11)
        pdf.set_text_color(26, 26, 26)
        pdf.set_x(20)
        display_content = _reshape_arabic(msg.content) if is_rtl else msg.content
        pdf.multi_cell(170, 6, display_content, new_x="LMARGIN", new_y="NEXT", align=align)
        pdf.ln(5)

    buf = io.BytesIO()
    pdf.output(buf)
    buf.seek(0)
    return buf


# ═══════════════════════════════════════════════════════════
#  DOCX — Message unique
# ═══════════════════════════════════════════════════════════

def _set_paragraph_rtl(paragraph, is_rtl: bool):
    """Définit la direction RTL d'un paragraphe DOCX (propriété w:bidi)."""
    from docx.oxml.ns import qn
    from lxml import etree
    pPr = paragraph._element.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if is_rtl:
        if bidi is None:
            bidi = etree.SubElement(pPr, qn("w:bidi"))
    else:
        if bidi is not None:
            pPr.remove(bidi)


def export_message_docx(message_content: str, message_role: str = "assistant", title: str = "") -> io.BytesIO:
    """Génère un DOCX pour un message unique avec RTL/LTR et w:bidi."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    lang = detect_language(message_content)
    is_rtl = lang == "ar"

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("EDUAI Learning")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
    run.font.italic = True

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("Export de reponse IA")
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph("─" * 60)

    display_title = title or "Reponse IA"
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(display_title)
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
    title_run.font.bold = True

    role_label = "Reponse de l'Assistant IA" if message_role == "assistant" else "Message de l'eleve"
    role_p = doc.add_paragraph()
    role_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role_run = role_p.add_run(f"● {role_label}")
    role_run.font.size = Pt(10)
    role_run.font.bold = True
    role_run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x2A)

    doc.add_paragraph("")

    content_p = doc.add_paragraph()
    _set_paragraph_rtl(content_p, is_rtl)
    if is_rtl:
        content_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        content_p.paragraph_format.right_indent = Inches(0.3)
    else:
        content_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        content_p.paragraph_format.left_indent = Inches(0.3)

    content_p.space_after = Pt(8)
    content_run = content_p.add_run(message_content)
    content_run.font.size = Pt(12)
    if is_rtl:
        content_run.font.name = "Tahoma"
        from docx.oxml.ns import qn
        rPr = content_run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            from lxml import etree
            rFonts = etree.SubElement(rPr, qn("w:rFonts"))
        rFonts.set(qn("w:ascii"), "Tahoma")
        rFonts.set(qn("w:hAnsi"), "Tahoma")
        rFonts.set(qn("w:eastAsia"), "Tahoma")

    doc.add_paragraph("")
    doc.add_paragraph("─" * 60)
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(
        f"EDUAI Learning — Plateforme educative intelligente — "
        f"Exporte le {_format_date(datetime.now())}"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ═══════════════════════════════════════════════════════════
#  DOCX — Conversation complète
# ═══════════════════════════════════════════════════════════

def export_docx(conversation, messages) -> io.BytesIO:
    """Génère un DOCX de conversation avec RTL/LTR par message et w:bidi."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("EDUAI Learning")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
    run.font.italic = True

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run("Export de conversation")
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph("─" * 60)

    title_p = doc.add_paragraph()
    title_run = title_p.add_run(conversation.title)
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
    title_run.font.bold = True

    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    meta_run = meta_p.add_run(
        f"Date : {_format_date(conversation.created_at)}  |  "
        f"Messages : {len(messages)}"
    )
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_paragraph("")

    orange_rgb = RGBColor(0xFF, 0x6B, 0x2B)
    navy_rgb = RGBColor(0x0D, 0x1B, 0x2A)

    for msg in messages:
        is_user = msg.role == "user"
        role_label = "Eleve" if is_user else "Assistant IA"
        color = orange_rgb if is_user else navy_rgb

        role_p = doc.add_paragraph()
        role_p.space_before = Pt(12)
        role_p.space_after = Pt(2)
        role_run = role_p.add_run(f"● {role_label}")
        role_run.font.size = Pt(10)
        role_run.font.bold = True
        role_run.font.color.rgb = color

        msg_lang = getattr(msg, "detected_language", None) or detect_language(msg.content)
        is_rtl = msg_lang == "ar"

        content_p = doc.add_paragraph()
        _set_paragraph_rtl(content_p, is_rtl)
        if is_rtl:
            content_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            content_p.paragraph_format.right_indent = Inches(0.3)
        else:
            content_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            content_p.paragraph_format.left_indent = Inches(0.3)

        content_p.space_after = Pt(8)
        content_run = content_p.add_run(msg.content)
        content_run.font.size = Pt(11)
        if is_rtl:
            content_run.font.name = "Tahoma"
            from docx.oxml.ns import qn
            rPr = content_run._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                from lxml import etree
                rFonts = etree.SubElement(rPr, qn("w:rFonts"))
            rFonts.set(qn("w:ascii"), "Tahoma")
            rFonts.set(qn("w:hAnsi"), "Tahoma")
            rFonts.set(qn("w:eastAsia"), "Tahoma")

    doc.add_paragraph("")
    doc.add_paragraph("─" * 60)
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(
        f"EDUAI Learning — Plateforme educative intelligente — "
        f"Exporte le {_format_date(datetime.now())}"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
